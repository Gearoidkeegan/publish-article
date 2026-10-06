#!/usr/bin/env python
"""Dump a .docx draft as styled text + emphasis + tables + text boxes + images.

Why a helper: OneDrive-backed .docx paths often fail to open in place
(python-docx raises PackageNotFoundError on the cloud-placeholder path). Copying
the file to a local temp first is the reliable workaround, so this script always
does that.

Usage:
    python extract-docx.py "<path to draft.docx>" [images-out-dir]

What it prints, in document order:
  [Style] text        one line per paragraph, with **bold** / *italic* preserved
  [Image] …           inline marker where a picture sits in the flow
  [TextBox n] …       author call-out boxes, including any table inside them
  [SmartArt n] …      diagram text, indented by its level in the diagram's tree
  --- table n         every table, pipe-joined, body and in-text-box alike
  === CONTENT CENSUS  counts of everything found, so nothing is silently dropped

WHY THE EXTRA CONTENT MATTERS. A draft's meaning does not live in plain
paragraphs alone. Four things carry author intent and were being dropped:

  * TEXT BOXES. Word stores call-outs, examples and side tables inside
    <w:txbxContent>, which `Document.paragraphs` and `Document.tables` do NOT
    traverse — a draft with a whole table of worked examples reported
    "TABLES: 0" and published without it. Text boxes are deliberate author
    emphasis: they must become a highlighted section in the article (a
    definition list, a table, or its own `##` section), never be silently
    dropped.
  * SMARTART DIAGRAMS. A SmartArt graphic keeps its text in a separate part,
    `word/diagrams/dataN.xml`, referenced from the paragraph by <dgm:relIds>.
    It is not in document.xml, not a <w:txbxContent>, and not an image part, so
    every other census line reports zero and the diagram vanishes without a
    trace. Often all that survives in a plain dump is a couple of loose arrow
    labels beside it, which read like stray call-outs. Publish the diagram's
    content as a list, a table, or a redrawn figure.
  * BOLD / ITALIC RUNS. Where the author bolded a defined term ("a **Repo** has
    two sides"), that emphasis is part of the writing and must survive into the
    markdown.
  * FOOTNOTES, ENDNOTES, HEADERS/FOOTERS. Rarer, but counted so you notice.

Check the CONTENT CENSUS before writing the markdown. If a count is non-zero,
that content belongs in the published article.
"""
import os
import sys
import shutil
import tempfile
import subprocess

try:
    import docx  # python-docx
    from docx.text.paragraph import Paragraph
    from docx.table import Table
except ImportError:
    sys.exit("python-docx not installed. Run: python -m pip install python-docx")

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
DGM = "{http://schemas.openxmlformats.org/drawingml/2006/diagram}"


def stage_copy(src, dst):
    """Copy src -> dst, resilient to OneDrive's Windows file-sharing lock.

    Python's open()/shutil.copyfile is often denied on OneDrive-backed files
    (PermissionError) while the shell `cp`/`copy` succeeds — different sharing
    mode. So try shutil first, then fall back to the shell copiers.
    """
    try:
        shutil.copyfile(src, dst)
        return
    except (PermissionError, OSError):
        pass
    for cmd in (["cp", src, dst], ["cmd", "/c", "copy", "/y", src, dst]):
        try:
            subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL,
                           stderr=subprocess.DEVNULL)
            if os.path.exists(dst) and os.path.getsize(dst) > 0:
                return
        except Exception:
            continue
    sys.exit("Could not read the .docx (OneDrive lock). Close it in Word and make "
             "sure it's downloaded (not cloud-only), then retry — or copy it to a "
             "local folder and pass that path.")


class Shim(object):
    """Minimal parent so Paragraph/Table proxies work on any element."""

    def __init__(self, part):
        self.part = part


def in_fallback(el):
    """True if el sits under <mc:Fallback> (the duplicate copy of a text box).

    Word writes every shape twice: <mc:Choice> for modern readers and
    <mc:Fallback> (VML) for old ones. Without this check every text box is
    reported twice.
    """
    node = el.getparent()
    while node is not None:
        if node.tag == MC + "Fallback":
            return True
        node = node.getparent()
    return False


def nested_in_textbox(el, root):
    """True if el sits inside a text box that is itself inside root.

    A text box is anchored *within* a paragraph, so p_el.iter() descends into
    it. Without this guard the whole call-out is glued onto its host
    paragraph's text, and a full section shows up as one giant run-on
    heading.
    """
    node = el.getparent()
    while node is not None and node is not root:
        if node.tag == W + "txbxContent":
            return True
        node = node.getparent()
    return False


def run_segments(p_el):
    """[(text, bold, italic)] for every run in p_el, in document order.

    Walks <w:r> directly rather than Paragraph.runs so runs nested inside
    <w:hyperlink> are included too.
    """
    segs = []
    for r in p_el.iter(W + "r"):
        if in_fallback(r) or nested_in_textbox(r, p_el):
            continue
        text = []
        for child in r:
            if child.tag == W + "t":
                text.append(child.text or "")
            elif child.tag == W + "tab":
                text.append("\t")
            elif child.tag in (W + "br", W + "cr"):
                text.append(" ")
        if not text:
            continue
        rPr = r.find(W + "rPr")
        bold = italic = False
        if rPr is not None:
            b, i = rPr.find(W + "b"), rPr.find(W + "i")
            bold = b is not None and b.get(W + "val") not in ("0", "false")
            italic = i is not None and i.get(W + "val") not in ("0", "false")
        segs.append(("".join(text), bold, italic))
    return segs


def render_text(p_el):
    """Paragraph text with **bold** / *italic* markers, adjacent runs merged."""
    merged = []
    for text, bold, italic in run_segments(p_el):
        if merged and merged[-1][1] == bold and merged[-1][2] == italic:
            merged[-1][0] += text
        else:
            merged.append([text, bold, italic])
    out = []
    for text, bold, italic in merged:
        if not text.strip():
            out.append(text)
            continue
        lead = text[:len(text) - len(text.lstrip())]
        trail = text[len(text.rstrip()):]
        core = text.strip()
        if bold:
            core = "**%s**" % core
        if italic:
            core = "*%s*" % core
        out.append(lead + core + trail)
    return "".join(out)


def para_style(p_el, shim):
    try:
        return Paragraph(p_el, shim).style.name
    except Exception:
        pPr = p_el.find(W + "pPr")
        if pPr is not None:
            st = pPr.find(W + "pStyle")
            if st is not None:
                return st.get(W + "val") or "Normal"
        return "Normal"


def image_refs(el, part):
    """Filenames of images anchored in el, in document order."""
    refs = []
    for blip in el.iter(A + "blip"):
        if in_fallback(blip) or nested_in_textbox(blip, el):
            continue
        rId = blip.get(R + "embed") or blip.get(R + "link")
        if not rId:
            continue
        try:
            refs.append(part.rels[rId].target_ref.split("/")[-1])
        except KeyError:
            refs.append(rId)
    return refs


def dgm_text(t_el):
    """Lines of a <dgm:t> body, with the author's **bold** / *italic* kept.

    SmartArt uses DrawingML text (<a:p>/<a:r>/<a:t>), not WordprocessingML, so
    render_text() cannot be reused here.
    """
    lines = []
    for p in t_el.findall(A + "p"):
        parts = []
        for r in p.iter(A + "r"):
            txt = "".join(t.text or "" for t in r.findall(A + "t"))
            if not txt.strip():
                parts.append(txt)
                continue
            rPr = r.find(A + "rPr")
            bold = rPr is not None and rPr.get("b") in ("1", "true")
            italic = rPr is not None and rPr.get("i") in ("1", "true")
            core = txt.strip()
            if bold:
                core = "**%s**" % core
            if italic:
                core = "*%s*" % core
            lead = txt[:len(txt) - len(txt.lstrip())]
            parts.append(lead + core + txt[len(txt.rstrip()):])
        line = "".join(parts).strip()
        if line:
            lines.append(line)
    return lines


def diagram_points(root):
    """[(lines, depth, kind)] for every authored point of a SmartArt data part.

    <dgm:ptLst> holds the nodes; <dgm:cxn type="parOf"> holds the tree, with
    srcOrd giving sibling order — so a nested diagram (a hierarchy, a funnel of
    funnels) comes out indented by level rather than flattened. Points of type
    "pres" are layout-only and carry no author text; "doc" is the root.
    """
    pts = {}
    for pt in root.iter(DGM + "pt"):
        pts[pt.get("modelId")] = pt
    kids = {}
    for cxn in root.iter(DGM + "cxn"):
        if cxn.get("type") != "parOf":
            continue
        try:
            order = int(cxn.get("srcOrd") or 0)
        except (TypeError, ValueError):
            order = 0
        kids.setdefault(cxn.get("srcId"), []).append((order, cxn.get("destId")))

    out, seen = [], set()

    def take(mid, depth):
        pt = pts[mid]
        kind = pt.get("type")
        if kind in ("pres", "doc"):
            return
        t = pt.find(DGM + "t")
        lines = dgm_text(t) if t is not None else []
        if lines:
            out.append((lines, max(depth - 1, 0), kind))

    def visit(mid, depth):
        if mid in seen or mid not in pts:
            return
        seen.add(mid)
        take(mid, depth)
        for _, child in sorted(kids.get(mid, [])):
            visit(child, depth + 1)

    for mid, pt in pts.items():
        if pt.get("type") == "doc":
            visit(mid, 0)
    # a point the tree never reached still holds author text: never drop it
    for mid, pt in pts.items():
        if mid not in seen:
            seen.add(mid)
            take(mid, 1)
    return out


def print_diagram(part, counters, indent=""):
    """Dump one SmartArt graphic's text, resolved from its diagramData part."""
    from lxml import etree                              # python-docx's own parser
    try:
        points = diagram_points(etree.fromstring(part.blob))
    except Exception as exc:                            # malformed / unknown part
        print("%s[SmartArt] could not read the diagram data: %s" % (indent, exc))
        return
    if not points:
        return
    counters["diagrams"] += 1
    n = counters["diagrams"]
    print("%s[SmartArt %d] %s" % (indent, n, "-" * 40))
    print("%s  ^ DIAGRAM TEXT — lives only in word/diagrams/*.xml, never in the "
          "paragraph flow. Author content: publish it as a list, a table or a "
          "redrawn figure." % indent)
    for lines, depth, kind in points:
        pad = indent + "    " + "  " * depth
        tag = "" if kind in (None, "node") else "   (%s)" % kind
        print("%s%s%s" % (pad, lines[0], tag))
        for extra in lines[1:]:
            print("%s%s" % (pad, extra))
    print("%s[/SmartArt %d]" % (indent, n))


def print_table(tbl_el, shim, counters, indent="", label=None):
    counters["tables"] += 1
    n = counters["tables"] - 1
    print("%s--- table %d%s" % (indent, n, (" (%s)" % label) if label else ""))
    try:
        rows = Table(tbl_el, shim).rows
        for row in rows:
            print(indent + " | ".join(c.text.replace("\n", " ") for c in row.cells))
        return
    except Exception:
        pass
    for tr in tbl_el.findall(W + "tr"):          # manual fallback
        cells = []
        for tc in tr.findall(W + "tc"):
            cells.append(" ".join(render_text(p) for p in tc.findall(W + "p")).strip())
        print(indent + " | ".join(cells))


def walk(container, shim, counters, indent=""):
    """Emit paragraphs / tables / text boxes of a container, in document order."""
    for el in container:
        if el.tag == W + "p":
            if in_fallback(el):
                continue
            counters["paragraphs"] += 1
            txt = render_text(el)
            if txt.strip():
                print("%s[%s] %s" % (indent, para_style(el, shim), txt))
            for ref in image_refs(el, shim.part):
                counters["images_inline"] += 1
                print("%s[Image] %s  <-- picture anchored here" % (indent, ref))
            # a SmartArt graphic points at its own data part from here
            for rel_ids in el.iter(DGM + "relIds"):
                if in_fallback(rel_ids) or nested_in_textbox(rel_ids, el):
                    continue
                rid = rel_ids.get(R + "dm")
                try:
                    part = shim.part.rels[rid].target_part
                except Exception:
                    counters["diagrams"] += 1
                    print("%s[SmartArt] unresolved diagram data (%s) — open the "
                          "draft and transcribe the figure by hand." % (indent, rid))
                    continue
                print_diagram(part, counters, indent)
            # text boxes are anchored inside a paragraph's drawing/pict
            for tb in el.iter(W + "txbxContent"):
                if in_fallback(tb):
                    continue
                counters["textboxes"] += 1
                n = counters["textboxes"]
                print("%s[TextBox %d] %s" % (indent, n, "-" * 40))
                print("%s  ^ AUTHOR CALL-OUT — must appear in the article, "
                      "highlighted (own section, definition list or table)." % indent)
                walk(tb, shim, counters, indent + "    ")
                print("%s[/TextBox %d]" % (indent, n))
        elif el.tag == W + "tbl":
            if in_fallback(el):
                continue
            print_table(el, shim, counters, indent,
                        "inside text box" if indent else None)


def part_text(part):
    """Non-empty paragraph texts in an arbitrary part (notes, headers, footers)."""
    try:
        root = part.element
    except AttributeError:
        return []
    return [render_text(p) for p in root.iter(W + "p")
            if render_text(p).strip() and not in_fallback(p)]


def extract_images(d, out_dir):
    """Write every embedded image part to out_dir. Returns [(filename, bytes_len)]."""
    saved = []
    os.makedirs(out_dir, exist_ok=True)
    for rel in d.part.rels.values():
        if "image" not in rel.reltype:
            continue
        part = rel.target_part
        name = part.partname.rpartition("/")[2]  # e.g. image1.png
        blob = part.blob
        with open(os.path.join(out_dir, name), "wb") as fh:
            fh.write(blob)
        saved.append((name, len(blob)))
    return saved


def main():
    # Windows consoles default to cp1252 and crash on characters such as
    # zero-width spaces that Word drafts often carry.
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except AttributeError:
            pass
    if len(sys.argv) < 2:
        sys.exit("usage: python extract-docx.py <path to draft.docx> [images-out-dir]")
    src = sys.argv[1]
    if not os.path.exists(src):
        sys.exit("not found: %s" % src)
    img_dir = sys.argv[2] if len(sys.argv) > 2 else os.path.join(
        tempfile.gettempdir(), "publish-article-images")

    # OneDrive workaround: copy to a local temp before opening.
    tmp = os.path.join(tempfile.gettempdir(), "publish-article-draft.docx")
    stage_copy(src, tmp)
    d = docx.Document(tmp)
    shim = Shim(d.part)

    counters = {"paragraphs": 0, "tables": 0, "textboxes": 0, "images_inline": 0,
                "diagrams": 0}
    walk(d.element.body, shim, counters)

    # parts python-docx's Document.paragraphs never reaches
    extras = []
    for rel in d.part.rels.values():
        rt = rel.reltype.rsplit("/", 1)[-1]
        if rt not in ("footnotes", "endnotes", "header", "footer", "comments"):
            continue
        lines = [t for t in part_text(rel.target_part) if t.strip()]
        if lines:
            extras.append((rt, lines))
    for rt, lines in extras:
        print("=== %s" % rt.upper())
        for t in lines:
            print("    %s" % t)

    saved = extract_images(d, img_dir)

    print("=== CONTENT CENSUS")
    print("    paragraphs   : %d" % counters["paragraphs"])
    print("    tables       : %d   (body + text boxes)" % counters["tables"])
    print("    text boxes   : %d" % counters["textboxes"])
    print("    smartart     : %d   (diagrams; text is in word/diagrams/, not the flow)"
          % counters["diagrams"])
    print("    images       : %d" % len(saved))
    print("    notes/headers: %d part(s) with text" % len(extras))
    if counters["diagrams"]:
        print("    !! SMARTART FOUND. A diagram is author content and counts as "
              "NEITHER an image nor a text box, so no other line here reports it. "
              "Publish each one as a list, a table or a redrawn figure — a draft "
              "whose figure is only a diagram will otherwise read complete while "
              "missing its whole point.")
    if counters["textboxes"]:
        print("    !! TEXT BOXES FOUND. These are author call-outs (examples, case "
              "studies, side tables). They MUST appear in the published article, "
              "highlighted — own '## ' section, a table, or a definition list. Do not drop them.")
    if counters["tables"]:
        print("    !! TABLES FOUND. Reproduce each as a markdown pipe table.")
    print("=== IMAGES: %d (extracted to %s)" % (len(saved), img_dir))
    for name, n in saved:
        print("    %s  (%d KB)" % (os.path.join(img_dir, name), n // 1024))
    if saved:
        print("    ^ These are part of the article. Optimise (WebP q90 is ~10x "
              "smaller than PNG for gradient+text heroes), place into "
              "the site's image folder (see the style guide), and reference it from "
              "the markdown. "
              "Do NOT publish text-only.")
    print("=== NOTE: **bold** / *italic* above is the author's own emphasis. "
          "Carry it into the markdown.")


if __name__ == "__main__":
    main()
